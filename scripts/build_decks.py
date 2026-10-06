"""
python-pptx deck builder for the weekly political-sentiment reports.

Produces TWO decks per person from the same dataset JSON:
  1. <slug>_x_report.pptx   -> X / social listening
  2. <slug>_news_report.pptx -> print / digital news

Branded title slide (dark hero + party accent), then content slides that drop
in the matplotlib charts and write the insight bullets. Layout uses explicit
EMU positions tuned to avoid chart/text collisions.
"""
import os
import json
from pptx import Presentation
from pptx.util import Inches, Pt, Emu
from pptx.dml.color import RGBColor
from pptx.enum.text import PP_ALIGN, MSO_ANCHOR

W, H = Inches(13.333), Inches(7.5)


def hexc(h):
    return RGBColor.from_string(h.lstrip("#"))


def _txt(slide, l, t, w, h, text, size, color, bold=False,
         align=PP_ALIGN.LEFT, anchor=MSO_ANCHOR.TOP, font="Calibri"):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    tf.vertical_anchor = anchor
    p = tf.paragraphs[0]
    p.alignment = align
    r = p.add_run()
    r.text = text
    r.font.size = Pt(size)
    r.font.bold = bold
    r.font.color.rgb = color
    r.font.name = font
    return box


def _rect(slide, l, t, w, h, color):
    from pptx.enum.shapes import MSO_SHAPE
    shp = slide.shapes.add_shape(MSO_SHAPE.RECTANGLE, l, t, w, h)
    shp.fill.solid()
    shp.fill.fore_color.rgb = color
    shp.line.fill.background()
    shp.shadow.inherit = False
    return shp


def _bullets(slide, l, t, w, h, items, size, color, accent):
    box = slide.shapes.add_textbox(l, t, w, h)
    tf = box.text_frame
    tf.word_wrap = True
    for i, it in enumerate(items):
        p = tf.paragraphs[0] if i == 0 else tf.add_paragraph()
        p.space_after = Pt(8)
        r = p.add_run()
        r.text = "•  " + it
        r.font.size = Pt(size)
        r.font.color.rgb = color
        r.font.name = "Calibri"
    return box


def _title_slide(prs, data, subtitle):
    s = prs.slides.add_slide(prs.slide_layouts[6])
    b = data["branding"]
    _rect(s, 0, 0, W, H, hexc(b["dark"]))
    _rect(s, 0, Inches(5.05), W, Inches(0.12), hexc(b["accent"]))
    p = data["person"]
    _txt(s, Inches(0.9), Inches(1.5), Inches(11.5), Inches(0.6),
         p["party"] + "  |  WEEKLY MEDIA & SOCIAL SENTIMENT",
         18, hexc(b["accent"]), bold=True)
    _txt(s, Inches(0.9), Inches(2.2), Inches(11.5), Inches(1.4),
         p["name"], 54, RGBColor(0xFF, 0xFF, 0xFF), bold=True)
    _txt(s, Inches(0.9), Inches(3.5), Inches(11.5), Inches(0.6),
         p["role"], 18, RGBColor(0xCF, 0xCF, 0xCF))
    _txt(s, Inches(0.9), Inches(5.3), Inches(11.5), Inches(0.6),
         subtitle, 20, RGBColor(0xFF, 0xFF, 0xFF), bold=True)
    _txt(s, Inches(0.9), Inches(6.0), Inches(11.5), Inches(0.5),
         "Week of " + data["week"]["label"], 15,
         RGBColor(0xAA, 0xAA, 0xAA))
    return s


def _content_header(s, data, kicker, title):
    b = data["branding"]
    _rect(s, 0, 0, W, Inches(1.15), hexc(b["dark"]))
    _rect(s, 0, Inches(1.15), W, Inches(0.06), hexc(b["accent"]))
    _txt(s, Inches(0.7), Inches(0.18), Inches(11), Inches(0.4),
         kicker, 13, hexc(b["accent"]), bold=True)
    _txt(s, Inches(0.7), Inches(0.5), Inches(12), Inches(0.6),
         title, 26, RGBColor(0xFF, 0xFF, 0xFF), bold=True)


def _stat_card(s, l, t, w, value, label, accent):
    _rect(s, l, t, w, Inches(1.5), RGBColor(0xF4, 0xF4, 0xF6))
    _rect(s, l, t, Inches(0.08), Inches(1.5), accent)
    _txt(s, l + Inches(0.2), t + Inches(0.18), w - Inches(0.3), Inches(0.7),
         value, 30, accent, bold=True)
    _txt(s, l + Inches(0.2), t + Inches(0.92), w - Inches(0.3), Inches(0.5),
         label, 12, RGBColor(0x55, 0x55, 0x55))


def _pic(s, path, l, t, w):
    if os.path.exists(path):
        return s.shapes.add_picture(path, l, t, width=w)
    return None


def build_x_deck(data, charts_dir, out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    b = data["branding"]
    accent = hexc(b["accent"])
    x = data["x_data"]
    blank = prs.slide_layouts[6]

    _title_slide(prs, data, "X / Social Listening Report")

    # Overview + stat cards
    s = prs.slides.add_slide(blank)
    _content_header(s, data, "OVERVIEW", "The week on X at a glance")
    total = sum(x["sentiment"].values())
    cards = [
        (str(x.get("official_posts", 0)), "Posts from official handle"),
        (str(x.get("sample_size", total)), "Conversation items sampled"),
        (f'{x["sentiment"]["negative"]*100//(total or 1)}%', "Negative share"),
        (f'{x["sentiment"]["positive"]*100//(total or 1)}%', "Positive share"),
    ]
    cw = Inches(2.95)
    for i, (v, l) in enumerate(cards):
        _stat_card(s, Inches(0.7) + i * (cw + Inches(0.18)), Inches(1.55),
                   cw, v, l, accent)
    _bullets(s, Inches(0.7), Inches(3.35), Inches(12), Inches(3.5),
             data.get("insights", [])[:4], 16, hexc(b["ink"]), accent)

    # Sentiment + daily trend
    s = prs.slides.add_slide(blank)
    _content_header(s, data, "SENTIMENT", "Sentiment split & daily volume")
    _pic(s, os.path.join(charts_dir, "x_sentiment.png"),
         Inches(0.7), Inches(1.6), Inches(4.2))
    _pic(s, os.path.join(charts_dir, "x_daily.png"),
         Inches(5.2), Inches(1.9), Inches(7.6))

    # Issues
    s = prs.slides.add_slide(blank)
    _content_header(s, data, "ISSUES", "What drove negative conversation")
    _pic(s, os.path.join(charts_dir, "x_issues.png"),
         Inches(0.7), Inches(1.6), Inches(7.4))
    _bullets(s, Inches(8.5), Inches(1.7), Inches(4.3), Inches(5),
             x.get("issue_notes", [])[:5], 13, hexc(b["ink"]), accent)

    # Top posts
    s = prs.slides.add_slide(blank)
    _content_header(s, data, "TOP POSTS", "Highest-reach posts this week")
    t = Inches(1.55)
    for post in x.get("top_posts", [])[:4]:
        _rect(s, Inches(0.7), t, Inches(11.9), Inches(1.18),
              RGBColor(0xF4, 0xF4, 0xF6))
        _rect(s, Inches(0.7), t, Inches(0.08), Inches(1.18), accent)
        _txt(s, Inches(0.95), t + Inches(0.1), Inches(9.3), Inches(0.98),
             post["text"], 13, hexc(b["ink"]))
        _txt(s, Inches(10.4), t + Inches(0.32), Inches(2.0), Inches(0.6),
             f'{post.get("views", 0):,} views', 13, accent, bold=True,
             align=PP_ALIGN.RIGHT)
        t += Inches(1.32)

    prs.save(out_path)
    return out_path


def build_news_deck(data, charts_dir, out_path):
    prs = Presentation()
    prs.slide_width, prs.slide_height = W, H
    b = data["branding"]
    accent = hexc(b["accent"])
    accent2 = hexc(b["accent2"])
    n = data["news_data"]
    blank = prs.slide_layouts[6]

    _title_slide(prs, data, "Print / Digital News Report")

    # Overview
    s = prs.slides.add_slide(blank)
    _content_header(s, data, "OVERVIEW", "Press coverage at a glance")
    total = sum(n["sentiment"].values())
    cards = [
        (str(n.get("article_count", total)), "Articles tracked"),
        (str(len(n.get("publications", []))), "Publications"),
        (f'{n["sentiment"]["positive"]*100//(total or 1)}%', "Positive share"),
        (f'{n["sentiment"]["neutral"]*100//(total or 1)}%', "Neutral share"),
    ]
    cw = Inches(2.95)
    for i, (v, l) in enumerate(cards):
        _stat_card(s, Inches(0.7) + i * (cw + Inches(0.18)), Inches(1.55),
                   cw, v, l, accent)
    _bullets(s, Inches(0.7), Inches(3.35), Inches(12), Inches(3.5),
             n.get("notes", [])[:4], 16, hexc(b["ink"]), accent)

    # Sentiment + languages
    s = prs.slides.add_slide(blank)
    _content_header(s, data, "SENTIMENT", "Tone & language of coverage")
    _pic(s, os.path.join(charts_dir, "news_sentiment.png"),
         Inches(0.9), Inches(1.7), Inches(4.0))
    _pic(s, os.path.join(charts_dir, "news_languages.png"),
         Inches(7.2), Inches(1.7), Inches(4.0))

    # Topics
    s = prs.slides.add_slide(blank)
    _content_header(s, data, "TOPICS", "What the press covered")
    _pic(s, os.path.join(charts_dir, "news_topics.png"),
         Inches(0.7), Inches(1.6), Inches(7.6))
    _bullets(s, Inches(8.6), Inches(1.7), Inches(4.2), Inches(5),
             n.get("topic_notes", [])[:5], 13, hexc(b["ink"]), accent)

    # Publications
    s = prs.slides.add_slide(blank)
    _content_header(s, data, "PUBLICATIONS", "Where coverage appeared")
    _pic(s, os.path.join(charts_dir, "news_publications.png"),
         Inches(0.7), Inches(1.6), Inches(7.6))

    # Defining stories
    s = prs.slides.add_slide(blank)
    _content_header(s, data, "STORIES", "Defining stories of the week")
    t = Inches(1.55)
    for st in data.get("stories", [])[:5]:
        _rect(s, Inches(0.7), t, Inches(11.9), Inches(1.0),
              RGBColor(0xF4, 0xF4, 0xF6))
        _rect(s, Inches(0.7), t, Inches(0.08), Inches(1.0), accent2)
        _txt(s, Inches(0.95), t + Inches(0.08), Inches(11.4), Inches(0.4),
             st["title"], 14, accent, bold=True)
        _txt(s, Inches(0.95), t + Inches(0.46), Inches(11.4), Inches(0.5),
             st["detail"], 12, hexc(b["ink"]))
        t += Inches(1.12)

    prs.save(out_path)
    return out_path


if __name__ == "__main__":
    import sys
    data_path, slug, out_dir = sys.argv[1], sys.argv[2], sys.argv[3]
    with open(data_path) as f:
        data = json.load(f)
    charts = os.path.join(out_dir, "charts")
    build_x_deck(data, charts, os.path.join(out_dir, f"{slug}_x_report.pptx"))
    build_news_deck(data, charts,
                    os.path.join(out_dir, f"{slug}_news_report.pptx"))
    print("decks ->", out_dir)
